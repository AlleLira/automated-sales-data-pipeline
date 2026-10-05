"""Command-line entry point for the automated sales data pipeline."""

from __future__ import annotations

import logging
import sys
from datetime import datetime
from time import perf_counter

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.config import INPUT_DIR, LOG_DIR, LOG_FILE, OUTPUT_DIR, OUTPUT_FILE
from src.extract import discover_excel_files, read_sales_file
from src.transform import transform_valid_rows
from src.validate import validate_rows, validate_schema


def configure_logging() -> logging.Logger:
    """Configure console and file logs for one pipeline run."""
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("sales_pipeline")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()

    formatter = logging.Formatter(
        "%(asctime)s | %(levelname)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )
    file_handler = logging.FileHandler(LOG_FILE, mode="w", encoding="utf-8")
    file_handler.setFormatter(formatter)
    console_handler = logging.StreamHandler(sys.stdout)
    console_handler.setFormatter(logging.Formatter("%(message)s"))
    logger.addHandler(file_handler)
    logger.addHandler(console_handler)
    return logger


def _format_workbook(file_path) -> None:
    """Apply readable formatting to every output worksheet."""
    from openpyxl import load_workbook

    workbook = load_workbook(file_path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        for column_cells in worksheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            letter = get_column_letter(column_cells[0].column)
            worksheet.column_dimensions[letter].width = min(max_length + 2, 45)

    workbook.save(file_path)


def write_output(
    valid_rows: pd.DataFrame,
    rejected_rows: pd.DataFrame,
    summary: pd.DataFrame,
) -> None:
    """Write all pipeline results to one formatted Excel workbook."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        valid_rows.to_excel(writer, sheet_name="Valid Sales", index=False)
        rejected_rows.to_excel(writer, sheet_name="Rejected Rows", index=False)
        summary.to_excel(writer, sheet_name="Processing Summary", index=False)
    _format_workbook(OUTPUT_FILE)


def run_pipeline() -> int:
    """Execute extraction, validation, transformation, and loading."""
    logger = configure_logging()
    started = perf_counter()

    logger.info("=" * 60)
    logger.info("AUTOMATED SALES DATA PIPELINE")
    logger.info("=" * 60)

    files = discover_excel_files(INPUT_DIR)
    logger.info("[INFO] Files found: %d", len(files))
    if not files:
        logger.error("[ERROR] No .xlsx files found in %s", INPUT_DIR)
        return 1

    valid_frames: list[pd.DataFrame] = []
    rejected_frames: list[pd.DataFrame] = []
    summary_records: list[dict[str, object]] = []

    for file_path in files:
        logger.info("\nProcessing: %s", file_path.name)
        try:
            dataframe = read_sales_file(file_path)
        except Exception as exc:
            logger.exception("[ERROR] File could not be read: %s", exc)
            summary_records.append(
                {
                    "file": file_path.name,
                    "status": "READ ERROR",
                    "input_rows": 0,
                    "valid_rows": 0,
                    "rejected_rows": 0,
                    "details": str(exc),
                }
            )
            continue

        schema_ok, missing, unexpected = validate_schema(dataframe)
        if not schema_ok:
            details = f"missing={missing}; unexpected={unexpected}"
            logger.error("[ERROR] Invalid schema: %s", details)
            summary_records.append(
                {
                    "file": file_path.name,
                    "status": "SCHEMA ERROR",
                    "input_rows": len(dataframe),
                    "valid_rows": 0,
                    "rejected_rows": len(dataframe),
                    "details": details,
                }
            )
            schema_rejected = dataframe.copy()
            schema_rejected["rejection_reason"] = f"invalid schema: {details}"
            rejected_frames.append(schema_rejected)
            continue

        logger.info("[OK] Schema validated")
        result = validate_rows(dataframe)
        transformed = transform_valid_rows(result.valid_rows)
        valid_frames.append(transformed)
        rejected_frames.append(result.rejected_rows)

        logger.info("[OK] Data quality checks completed")
        logger.info("[OK] Categories standardized")
        logger.info("[OK] Sales totals calculated")
        if result.checks["rejected_rows"]:
            logger.warning(
                "[WARNING] Rejected rows: %d", result.checks["rejected_rows"]
            )

        summary_records.append(
            {
                "file": file_path.name,
                "status": "PROCESSED",
                **result.checks,
                "details": "",
            }
        )

    valid_sales = pd.concat(valid_frames, ignore_index=True) if valid_frames else pd.DataFrame()
    rejected_sales = (
        pd.concat(rejected_frames, ignore_index=True, sort=False)
        if rejected_frames
        else pd.DataFrame()
    )

    # Detect sale IDs repeated across different files after row-level validation.
    if not valid_sales.empty:
        duplicate_across_files = valid_sales["sale_id"].duplicated(keep=False)
        if duplicate_across_files.any():
            cross_file_rejected = valid_sales.loc[duplicate_across_files].copy()
            cross_file_rejected["rejection_reason"] = "duplicate sale_id across files"
            rejected_sales = pd.concat(
                [rejected_sales, cross_file_rejected], ignore_index=True, sort=False
            )
            valid_sales = valid_sales.loc[~duplicate_across_files].copy()

    summary = pd.DataFrame(summary_records)
    elapsed = perf_counter() - started
    totals = pd.DataFrame(
        [
            {
                "file": "TOTAL",
                "status": "COMPLETED",
                "input_rows": int(summary.get("input_rows", pd.Series(dtype=int)).sum()),
                "valid_rows": len(valid_sales),
                "rejected_rows": len(rejected_sales),
                "details": f"execution_time_seconds={elapsed:.3f}; run_at={datetime.now():%Y-%m-%d %H:%M:%S}",
            }
        ]
    )
    summary = pd.concat([summary, totals], ignore_index=True)

    write_output(valid_sales, rejected_sales, summary)

    logger.info("\n%s", "=" * 60)
    logger.info("PROCESS COMPLETED")
    logger.info("=" * 60)
    logger.info("Files processed: %d", len(files))
    logger.info("Valid rows: %d", len(valid_sales))
    logger.info("Rejected rows: %d", len(rejected_sales))
    logger.info("Execution time: %.3f seconds", elapsed)
    logger.info("Output: %s", OUTPUT_FILE)
    return 0


if __name__ == "__main__":
    raise SystemExit(run_pipeline())

